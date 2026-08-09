from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(
    "/Users/hankchen/Library/CloudStorage/OneDrive-SCIE/backup/Dev/Apps/"
    "Done&Uploaded/Floatem"
)
OUTPUT = ROOT / "output" / "email response" / "Apple开发者支持回复-TestFlight安装失败.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial Unicode MS", size=9.6, color="24292F", bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_body_paragraph(doc, text="", space_after=3, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.17
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=9.6)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.68)
    p.paragraph_format.first_line_indent = Cm(-0.42)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.space_after = Pt(1.5)
    set_run_font(p.add_run(text), size=9.2)
    return p


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2.5)
    run = p.add_run(text)
    set_run_font(run, size=10.7, color="0A66C2", bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.25)
section.bottom_margin = Cm(1.2)
section.left_margin = Cm(1.65)
section.right_margin = Cm(1.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial Unicode MS"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
normal.font.size = Pt(9.6)
normal.paragraph_format.space_after = Pt(3)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
set_run_font(
    title.add_run("TestFlight macOS 外部测试安装失败 - 补充排查信息"),
    size=14.5,
    color="111827",
    bold=True,
)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(8)
set_run_font(
    subtitle.add_run("Apple Developer Support 案例编号：20000119582501"),
    size=8.8,
    color="6B7280",
)

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
meta.columns[0].width = Cm(3.2)
meta.columns[1].width = Cm(13.3)
meta_data = [
    ("邮件主题", "关于 Floatem macOS 公开链接测试无法安装的补充排查结果"),
    ("App", "Floatem（Bundle ID：com.hankch.floatem）"),
    ("当前测试构建", "版本 1.0.8（Build 48），macOS 外部公开链接测试"),
]
for idx, (label, value) in enumerate(meta_data):
    left, right = meta.rows[idx].cells
    left.width = Cm(3.2)
    right.width = Cm(13.3)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(left, "EEF4FB")
    set_cell_shading(right, "F8FAFC")
    set_cell_margins(left, top=65, bottom=65)
    set_cell_margins(right, top=65, bottom=65)
    lp = left.paragraphs[0]
    rp = right.paragraphs[0]
    lp.paragraph_format.space_after = Pt(0)
    rp.paragraph_format.space_after = Pt(0)
    set_run_font(lp.add_run(label), size=8.8, color="334155", bold=True)
    set_run_font(rp.add_run(value), size=8.8, color="334155")

doc.add_paragraph().paragraph_format.space_after = Pt(0)

salute = doc.add_paragraph()
salute.paragraph_format.space_after = Pt(4)
set_run_font(salute.add_run("Water 您好："), size=9.6)

add_body_paragraph(
    doc,
    "感谢您就 TestFlight 问题提供排查建议。我们已根据邮件逐项核对测试方式、"
    "Xcode 版本、签名资源、TestFlight 客户端及测试员状态。以下为 Floatem "
    "macOS 外部测试的完整情况和当前复现结果。",
)

add_section_heading(doc, "一、App 与构建情况")
for item in [
    "App 名称：Floatem；Bundle ID：com.hankch.floatem。",
    "当前公开测试构建：1.0.8（Build 48），发布日期为 2026 年 7 月 25 日。",
    "发布和管理此 App 所使用的是 Fengyan Zhang 的 Apple Developer / App Store Connect 账户。",
    "本次为 macOS 外部测试，采用 TestFlight 公开链接方式；公开测试页面要求 macOS 14 或更高版本。",
    "问题最早约在 1.0.7 测试阶段、上周日附近出现；目前 1.0.8（Build 48）仍可稳定复现。该首次出现时间为估算。",
]:
    add_bullet(doc, item)

add_section_heading(doc, "二、开发工具与签名资源")
for item in [
    "用于构建和上传的 Xcode 为 26.0.1（17A400），高于 macOS TestFlight 所要求的 Xcode 13。",
    "Xcode 账户为 Fengyan Zhang。Manage Certificates 页面目前显示 Apple Development、Apple Distribution 和 Developer ID Application 证书，相关截图已附上。",
    "Build 48 已成功上传至 App Store Connect、完成处理并进入外部测试；归档的主 App 已通过 codesign --verify --deep --strict 验证。",
    "归档支持 arm64 和 x86_64，最低系统版本为 macOS 14.0；签名中的 Team ID、Bundle ID 和 App Group 均一致。",
    "烦请同时确认：对于当前采用 Xcode Organizer / App Store Connect 的 TestFlight 上传流程，是否还必须单独配置本地 Mac Installer Distribution 证书，或可由 Apple 的云签名流程完成。",
]:
    add_bullet(doc, item)

add_section_heading(doc, "三、测试员与加入方式")
add_body_paragraph(
    doc,
    "App Store Connect“所有测试员”页面中，共显示三条记录。其中两条带邮箱的账户属于我们的 "
    "Development Team；本次公开链接外部测试组实际包含以下两名已接受测试的用户：",
)

tester_table = doc.add_table(rows=1, cols=4)
tester_table.alignment = WD_TABLE_ALIGNMENT.CENTER
tester_table.autofit = False
widths = [Cm(3.0), Cm(4.8), Cm(4.4), Cm(4.0)]
headers = ["测试员", "Apple ID / 账户", "加入方式", "结果"]
for i, (cell, text, width) in enumerate(zip(tester_table.rows[0].cells, headers, widths)):
    cell.width = width
    set_cell_shading(cell, "0A66C2")
    set_cell_margins(cell, top=55, bottom=55)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run(text), size=8.5, color="FFFFFF", bold=True)
set_repeat_table_header(tester_table.rows[0])

testers = [
    ("虞可馨", "19558013109\n（公开链接中显示为匿名）", "公开链接；2026-07-25 已接受", "无法安装"),
    ("Hank Chen / 陈涵", "hankchenchh@gmail.com", "公开链接；2026-07-25 已接受", "无法安装"),
]
for row_idx, row_data in enumerate(testers):
    row = tester_table.add_row()
    for i, (cell, text, width) in enumerate(zip(row.cells, row_data, widths)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
        set_cell_margins(cell, top=65, bottom=65)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 2, 3) else WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p.add_run(text), size=8.2, color="334155")

add_body_paragraph(
    doc,
    "说明：Hank Chen 的账户同时是 Development Team 成员，因此在“所有测试员”页面中会显示“内部”"
    "账户类型；但本次加入 Floatem test group 的操作是通过公开链接完成。更重要的是，非团队成员虞可馨"
    "以匿名公开链接测试员身份也出现完全相同的安装失败，因此问题并非仅发生在团队成员账户上。",
    first_line=True,
)

add_section_heading(doc, "四、TestFlight 客户端与失败现象")
for item in [
    "两名测试员均从 Mac App Store 下载或更新至当时可获得的最新版 TestFlight App。",
    "两名测试员使用的 Mac 均满足公开测试页面所示的 macOS 14 或更高版本要求。",
    "公开链接可以正常打开，Floatem 可以正常出现在 TestFlight 中，版本号、Build 号和“安装”按钮均可见。",
    "点击“安装”后，两名测试员均立即收到相同提示：“无法安装 App / 无法安装 Floatem。TestFlight 目前不可用。请重试。”",
    "错误可重复出现，卸载旧版本或重新打开 TestFlight 后仍未恢复。",
]:
    add_bullet(doc, item)

add_section_heading(doc, "五、已捕获的技术信息")
add_body_paragraph(
    doc,
    "在一次现场复现中，TestFlight 已成功发起 Build 48 的安装请求，但 Apple 的安装数据接口在下载"
    "任何 App 载荷之前返回 HTTP 500。日志显示 Error Downloading Install Data，且 "
    "downloadProgress 和 installProgress 均为 null。这与两名测试员看到的客户端提示一致。",
)

log_table = doc.add_table(rows=4, cols=2)
log_table.alignment = WD_TABLE_ALIGNMENT.CENTER
log_table.autofit = False
log_rows = [
    ("TestFlight Build ID", "224777408"),
    ("Request ID", "0BFFEBA5-9B0C-4EC1-B010-A3CF22A9C194"),
    ("Correlation Key", "2B4FNGP3NEP5B2XZC5MT466UMM"),
    ("失败阶段", "ProcessingInstallInitiateResponse / HTTP 500"),
]
for idx, (label, value) in enumerate(log_rows):
    left, right = log_table.rows[idx].cells
    left.width = Cm(4.1)
    right.width = Cm(12.4)
    set_cell_shading(left, "EEF4FB")
    set_cell_shading(right, "F8FAFC")
    set_cell_margins(left, top=70, bottom=70)
    set_cell_margins(right, top=70, bottom=70)
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    right.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_run_font(left.paragraphs[0].add_run(label), size=8.9, color="334155", bold=True)
    set_run_font(right.paragraphs[0].add_run(value), size=8.9, color="334155")

add_section_heading(doc, "六、请求协助")
add_body_paragraph(
    doc,
    "由于一个开发团队成员账户和一个完全独立的匿名外部测试账户都能看到构建、但都在获取安装数据"
    "阶段收到相同错误，烦请协助检查 Floatem 对应 TestFlight 构建的服务端安装数据、分发包生成状态"
    "或账号侧配置是否异常，并确认当前证书组合是否满足 macOS TestFlight 分发要求。如需我们重新上传"
    "诊断构建、提供完整日志或补充设备信息，请告知具体要求。",
)

add_section_heading(doc, "附件清单")
attachments = [
    "xcode版本.png - Xcode 26.0.1（17A400）版本页面",
    "xcode签名.png - Fengyan Zhang 账户的签名证书页面",
    "test group.png - Floatem 外部组、公开链接和两名已接受测试员",
    "所有测试员.png - App Store Connect 所有测试员列表",
    "testflight失败.png - 测试员一的完整 TestFlight 失败窗口",
    "testflight失败-2.png - 测试员二的完整 TestFlight 失败窗口",
]
for item in attachments:
    add_bullet(doc, item)

closing = doc.add_paragraph()
closing.paragraph_format.space_before = Pt(5)
closing.paragraph_format.space_after = Pt(3)
set_run_font(closing.add_run("感谢协助。"), size=9.6)

signature = doc.add_paragraph()
signature.paragraph_format.space_after = Pt(0)
set_run_font(signature.add_run("此致\nFengyan Zhang 开发团队"), size=9.6)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(
    footer.add_run("Apple Developer Support Case 20000119582501"),
    size=8,
    color="94A3B8",
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
